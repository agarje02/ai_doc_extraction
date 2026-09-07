"""Text extraction across formats with an OCR fallback for scanned content.

Strategy per format:
  - pdf   : native text via pdfplumber; if a page is (nearly) empty and OCR is
            enabled, render it to an image and run tesseract.
  - image : tesseract OCR.
  - docx  : python-docx paragraphs + tables.
  - text  : read directly.

Every extractor returns a list of ``PageText`` so downstream chunking can keep
page provenance for the UI.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from app.config import Settings, get_settings


@dataclass
class PageText:
    page_number: int
    text: str
    method: str  # "native" | "ocr" | "docx" | "text"


@dataclass
class ExtractedText:
    pages: list[PageText] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)

    @property
    def used_ocr(self) -> bool:
        return any(p.method == "ocr" for p in self.pages)


def extract_text(path: str | Path, kind: str, settings: Settings | None = None) -> ExtractedText:
    settings = settings or get_settings()
    path = Path(path)
    if kind == "pdf":
        return _extract_pdf(path, settings)
    if kind == "image":
        return ExtractedText(pages=[PageText(1, _ocr_image(path, settings), "ocr")])
    if kind == "docx":
        return ExtractedText(pages=[PageText(1, _extract_docx(path), "docx")])
    if kind == "text":
        return ExtractedText(
            pages=[PageText(1, path.read_text(encoding="utf-8", errors="replace"), "text")]
        )
    raise ValueError(f"Unsupported document kind: {kind!r}")


def _extract_pdf(path: Path, settings: Settings) -> ExtractedText:
    import pdfplumber

    pages: list[PageText] = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            native = (page.extract_text() or "").strip()
            if len(native) >= settings.ocr_min_chars or not settings.ocr_enabled:
                pages.append(PageText(i, native, "native"))
            else:
                ocr_text = _ocr_pdf_page(path, i, settings)
                if ocr_text.strip():
                    pages.append(PageText(i, ocr_text, "ocr"))
                else:
                    pages.append(PageText(i, native, "native"))
    return ExtractedText(pages=pages)


def _configure_tesseract(settings: Settings) -> None:
    import pytesseract

    if settings.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd


def _ocr_image(path: Path, settings: Settings) -> str:
    import pytesseract
    from PIL import Image

    _configure_tesseract(settings)
    with Image.open(path) as img:
        return pytesseract.image_to_string(img)


def _ocr_pdf_page(path: Path, page_number: int, settings: Settings) -> str:
    import pytesseract
    from pdf2image import convert_from_path

    _configure_tesseract(settings)
    kwargs = {"first_page": page_number, "last_page": page_number, "dpi": 200}
    if settings.poppler_path:
        kwargs["poppler_path"] = settings.poppler_path
    images = convert_from_path(str(path), **kwargs)
    return "\n".join(pytesseract.image_to_string(img) for img in images)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)
